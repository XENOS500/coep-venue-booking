"""Rebuild the traceable catalog fixture using the supplied workbook; no database writes.
Requires openpyxl (pip install openpyxl). Run from the repository root.
"""
import json
import re
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[2]
sheet = openpyxl.load_workbook(root / 'coep_venues.xlsx')['in']
records = []
for row in sheet.iter_rows(min_row=2, max_row=62):
    raw = {openpyxl.utils.get_column_letter(i + 1): cell.value for i, cell in enumerate(row) if cell.value is not None}
    clean = lambda i: str(row[i].value).strip() if row[i].value is not None else None
    image = row[11]
    url = image.hyperlink.target if image.hyperlink else None
    if not url and image.value:
        match = re.search(r'https://\S+', str(image.value))
        url = match.group(0) if match else None
    capacity = row[9].value
    capacity = int(capacity) if isinstance(capacity, (int, float)) and capacity > 0 else None
    notes = ['Verify opening hours and booking configuration.', 'Review source floor, venue type and room reference; workbook values may be ambiguous.']
    if capacity is None:
        notes.append('Capacity is missing; verification required before booking.')
    if url:
        notes.append('Image source is unverified; shared or folder links may need authorized asset selection.')
    elif clean(11):
        notes.append('Unresolved image label; no linked asset.')
    records.append(dict(sourceSerial=int(row[0].value), buildingName=clean(1), floor=clean(2), name=clean(3), venueType=clean(4), roomReferenceRaw=clean(5), capacity=capacity, features=clean(10), imageSourceUrl=url, imageSourceLabel=clean(11), allowedUses=clean(12), sourceRaw={**raw, 'imageHyperlink': image.hyperlink.target if image.hyperlink else None}, reviewNotes=notes))
assert len(records) == 61
assert {r['sourceSerial'] for r in records} == set(range(1, 63)) - {50}
target = root / 'server/prisma/data/venues.json'
target.parent.mkdir(parents=True, exist_ok=True)
target.write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
print(f'Extracted {len(records)} venues, {sum(bool(r["imageSourceUrl"]) for r in records)} linked sources.')
